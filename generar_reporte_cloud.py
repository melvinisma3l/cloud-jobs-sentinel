# ==============================================================================
# CLOUD AUTOMATION DÍA 16: PERSISTENCIA Y DESCARGA DE REPORTES (ARTIFACTS)
# OBJETIVO: Generar un Excel corporativo en la nube y dejarlo descargable en GitHub
# ==============================================================================

import sys
from datetime import datetime
import requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")

def generar_reporte_nube():
    ahora = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    print("=" * 70)
    print(f"🤖 GENERADOR CLOUD INICIADO | TIMESTAMP: {ahora}")
    print("=" * 70)

    url = "https://remoteok.com/api"
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) CloudWatcher/2.0"}

    try:
        resp = requests.get(url, headers=headers, timeout=15)
        if resp.status_code != 200:
            print(f"⚠️ Error HTTP: {resp.status_code}")
            return

        datos = resp.json()
        empleos = [item for item in datos if isinstance(item, dict) and "position" in item][:15]

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Remote Tech Intelligence"
        ws.views.sheetView[0].showGridLines = True

        NAVY = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        ZEBRA = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
        WHITE = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

        FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        FONT_TITLE  = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
        FONT_BOLD   = Font(name="Calibri", size=10, bold=True, color="0F172A")
        FONT_BODY   = Font(name="Calibri", size=10, color="334155")
        FONT_LINK   = Font(name="Calibri", size=10, color="2563EB", underline="single")

        THIN = Side(border_style="thin", color="CBD5E1")
        BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

        ws["A1"] = "GLOBAL REMOTE JOBS — DAILY CLOUD DISPATCH"
        ws["A1"].font = FONT_TITLE
        ws["A2"] = f"Automated Cloud Execution via GitHub Actions | Generated: {ahora} UTC"
        ws["A2"].font = Font(name="Calibri", size=10, color="64748B")

        headers = ["Job Title", "Company", "Tags / Stack", "Location", "Direct Apply URL"]
        for col_idx, text in enumerate(headers, start=1):
            c = ws.cell(row=4, column=col_idx, value=text)
            c.fill = NAVY
            c.font = FONT_HEADER
            c.alignment = Alignment(horizontal="center" if col_idx > 3 else "left", vertical="center")
            c.border = BORDER

        ws.row_dimensions[4].height = 26

        for idx, j in enumerate(empleos, start=5):
            puesto = j.get("position", "N/A")
            empresa = j.get("company", "N/A")
            tags = ", ".join(j.get("tags", [])[:4])
            ubicacion = j.get("location", "Worldwide") or "Worldwide"
            link = j.get("url", "")
            hyperlink = f'=HYPERLINK("{link}", "Apply Now ↗")' if link else "N/A"

            row_fill = ZEBRA if idx % 2 == 0 else WHITE

            ws.cell(row=idx, column=1, value=puesto).font = FONT_BOLD
            ws.cell(row=idx, column=2, value=empresa).font = FONT_BODY
            ws.cell(row=idx, column=3, value=tags).font = FONT_BODY
            ws.cell(row=idx, column=4, value=ubicacion).font = FONT_BODY
            ws.cell(row=idx, column=5, value=hyperlink).font = FONT_LINK

            for c_idx in range(1, 6):
                c = ws.cell(row=idx, column=c_idx)
                c.fill = row_fill
                c.border = BORDER
                c.alignment = Alignment(horizontal="center" if c_idx in [4, 5] else "left", vertical="center")

            ws.row_dimensions[idx].height = 20

        ws.column_dimensions["A"].width = 38
        ws.column_dimensions["B"].width = 22
        ws.column_dimensions["C"].width = 32
        ws.column_dimensions["D"].width = 18
        ws.column_dimensions["E"].width = 18

        nombre_archivo = "Daily_Remote_Jobs_Report.xlsx"
        wb.save(nombre_archivo)
        print(f"✅ Archivo Excel generado con éxito en la nube: {nombre_archivo}")
        print("=" * 70)

    except Exception as e:
        print(f"❌ Error durante la generación: {e}")

if __name__ == "__main__":
    generar_reporte_nube()
